"""Tests for the independent Ozon -> M.Video single-SKU dry-run workflow."""

from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text

import backend.app.single_sku.workflow as workflow
from backend.app.database import ensure_single_sku_jobs_columns
from backend.app.models import CategoryMapping, SingleSkuStatus
from backend.app.single_sku.pricing import calculate_price


def sample_payload() -> dict:
    return {
        "url": "https://www.ozon.ru/product/1237743084/",
        "product_id": "1237743084",
        "sku_id": "TQBSM1",
        "title": "Acme Станок для резки ленты TQBSM1",
        "description": "Acme диспенсер TQBSM1 для монтажной ленты.",
        "brand": "Acme",
        "model": "TQBSM1",
        "price_rub": 1999,
        "category": {"id": 17029021, "name": "Диспенсеры для монтажной ленты"},
        "packageInfo": {
            "weightG": 400,
            "lengthMm": 130,
            "widthMm": 65,
            "heightMm": 80,
        },
        "material": "Силикон",
        "images": [
            "https://example.com/main.jpg",
            "https://example.com/second.jpg",
        ],
    }


def confirm_category(session) -> None:
    session.add(
        CategoryMapping(
            ozon_category_id=17029021,
            ozon_category_name="Диспенсеры для монтажной ленты",
            mv_group_id="604171101",
            mv_group_name="Диспенсеры для монтажной ленты",
            mv_infomodel_id="INF-307590",
            status="confirmed",
        )
    )
    session.commit()


def create_confirmed_job(session):
    confirm_category(session)
    return workflow.create_job_from_ozon_payload(session, sample_payload())


def test_converts_ozon_units_exactly() -> None:
    converted = workflow.convert_package_dimensions(
        length_mm="130",
        width_mm="65",
        height_mm="80",
        weight_g="400",
    )

    assert converted["length_cm"] == Decimal("13")
    assert converted["width_cm"] == Decimal("6.5")
    assert converted["height_cm"] == Decimal("8")
    assert converted["weight_kg"] == Decimal("0.4")


def test_dry_run_prices_writes_excel_and_does_not_upload(memory_session) -> None:
    job = create_confirmed_job(memory_session)

    result = workflow.prepare_dry_run(
        memory_session,
        job,
        purchase_cost_cny=Decimal("12.34"),
        stock=5,
        color="фиолетовый",
    )

    assert result.status == SingleSkuStatus.DRY_RUN_READY
    assert result.source_price_rub == Decimal("1999.00")
    assert result.purchase_cost_cny == Decimal("12.3400")
    assert result.stock == 5
    assert result.upload_ref == ""
    assert result.uploaded_at is None

    expected_pricing = calculate_price(
        purchase_cost_cny=Decimal("12.34"),
        domestic_cost_cny=Decimal("0"),
        weight_g=Decimal("400"),
        length_cm=Decimal("13"),
        width_cm=Decimal("6.5"),
        height_cm=Decimal("8"),
        commission_category="home",
        target_net_margin=Decimal("0.25"),
        shipping_channel="economy",
    )
    assert result.price_rub == expected_pricing["price_rub"]

    assert "Acme" not in result.sanitized_title
    assert "Acme" not in result.sanitized_description
    assert "TQBSM1" in result.sanitized_title
    assert "TQBSM1" in result.sanitized_description

    dry_run = result.dry_run_result_json
    assert dry_run["upload_performed"] is False
    assert dry_run["upload_ref"] == ""
    assert dry_run["column_count"] == 95
    assert dry_run["source_price_rub_used_as_cny_cost"] is False

    workbook = load_workbook(dry_run["output_file_path"])
    worksheet = workbook[dry_run["worksheet_name"]]
    assert worksheet.max_column >= 95
    assert worksheet.cell(5, 1).value == "Сгенерировать"
    assert worksheet.cell(5, 2).value == result.sanitized_title
    assert worksheet.cell(5, 3).value == "Нет бренда"
    assert worksheet.cell(5, 7).value == "TQBSM1"
    assert worksheet.cell(5, 9).value == result.job_ref
    assert worksheet.cell(5, 14).value == 13.0
    assert worksheet.cell(5, 15).value == 6.5
    assert worksheet.cell(5, 16).value == 8.0
    assert worksheet.cell(5, 17).value == 0.4
    assert worksheet.cell(5, 48).value == 8.0
    assert worksheet.cell(5, 49).value == 6.5
    assert worksheet.cell(5, 50).value == 13.0
    assert worksheet.cell(5, 58).value == "Силикон"
    assert worksheet.cell(5, 64).value == 0.4
    assert worksheet.cell(5, 70).value == "Силикон"
    assert worksheet.cell(5, 76).value == "https://example.com/main.jpg"
    assert worksheet.cell(5, 77).value == "https://example.com/second.jpg"


def test_ozon_rub_is_not_automatically_used_as_cny_cost(memory_session) -> None:
    job = create_confirmed_job(memory_session)

    with pytest.raises(workflow.SingleSkuWorkflowError):
        workflow.prepare_dry_run(
            memory_session,
            job,
            purchase_cost_cny=None,
            stock=5,
        )

    assert job.status == SingleSkuStatus.BLOCKED
    assert job.purchase_cost_cny is None
    assert job.source_price_rub == Decimal("1999.00")
    assert "CNY" in job.last_error


def test_missing_confirmed_category_mapping_blocks(memory_session) -> None:
    job = workflow.create_job_from_ozon_payload(memory_session, sample_payload())

    with pytest.raises(workflow.SingleSkuWorkflowError):
        workflow.prepare_dry_run(
            memory_session,
            job,
            purchase_cost_cny=Decimal("12.34"),
            stock=5,
        )

    assert job.status == SingleSkuStatus.BLOCKED
    assert "confirmed" in job.last_error


def test_missing_stock_blocks(memory_session) -> None:
    job = create_confirmed_job(memory_session)

    with pytest.raises(workflow.SingleSkuWorkflowError):
        workflow.prepare_dry_run(
            memory_session,
            job,
            purchase_cost_cny=Decimal("12.34"),
            stock=0,
        )

    assert job.status == SingleSkuStatus.BLOCKED
    assert "库存" in job.last_error


def test_certificate_and_tn_ved_precheck_blocks_then_allows(memory_session, monkeypatch) -> None:
    strict_rule = workflow.CategoryPublishRule(
        ozon_category_id=17029021,
        mv_group_id="604171101",
        mv_infomodel_id="INF-307590",
        template_file_path="work/template_dispenser.xlsx",
        commission_category="home",
        tn_ved_required=True,
        required_certificates=("ДС",),
        brand_authorization_required=False,
        required_material=True,
    )
    monkeypatch.setattr(
        workflow,
        "CATEGORY_PUBLISH_RULES",
        {17029021: strict_rule},
    )
    job = create_confirmed_job(memory_session)

    with pytest.raises(workflow.SingleSkuWorkflowError):
        workflow.prepare_dry_run(
            memory_session,
            job,
            purchase_cost_cny=Decimal("12.34"),
            stock=5,
        )
    assert job.status == SingleSkuStatus.BLOCKED
    assert "ДС" in job.last_error
    assert "TN VED" in job.last_error

    result = workflow.prepare_dry_run(
        memory_session,
        job,
        purchase_cost_cny=Decimal("12.34"),
        stock=5,
        tn_ved="8466950000",
        compliance_documents=[
            {
                "cert_type": "ДС",
                "reg_number": "EAC-TEST-001",
                "file_url": "https://example.com/cert.pdf",
            }
        ],
    )
    assert result.status == SingleSkuStatus.DRY_RUN_READY
    assert result.tn_ved == "8466950000"


def test_adds_columns_to_historical_single_sku_table(tmp_path) -> None:
    db_path = tmp_path / "old-mvideo.db"
    engine = create_engine(
        f"sqlite:///{db_path}",
        connect_args={"check_same_thread": False},
    )

    with engine.begin() as connection:
        connection.execute(text("DROP TABLE IF EXISTS single_sku_jobs"))
        connection.execute(
            text(
                """
                CREATE TABLE single_sku_jobs (
                    id INTEGER PRIMARY KEY,
                    job_ref VARCHAR(64) NOT NULL,
                    source_platform VARCHAR(32) NOT NULL,
                    status VARCHAR(32) NOT NULL
                )
                """
            )
        )

    ensure_single_sku_jobs_columns(engine)

    columns = {column["name"] for column in inspect(engine).get_columns("single_sku_jobs")}
    assert {
        "template_file_path",
        "image_status",
        "upload_ref",
        "uploaded_at",
        "compliance_documents_json",
        "color",
    }.issubset(columns)
