# -*- coding: utf-8 -*-
"""Tests for app.pipeline.template_builder + image_pipeline (v0.3).

Pure row-building is tested without openpyxl; the xlsx write uses the real
blank M.Video template. OSS uploads use a recording fake uploader.
"""

from pathlib import Path

import openpyxl
from PIL import Image

from app.pipeline.image_pipeline import (
    process_local_images,
    upload_processed_to_oss,
)
from app.pipeline.template_builder import (
    COLS,
    DATA_START_ROW,
    build_template_rows,
    write_template_xlsx,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "work" / "mvideo_template.xlsx"


class FakeUploader:
    def __init__(self):
        self.calls: list[tuple[str, str]] = []

    def upload_image_file(self, local_path: str, object_key: str) -> str:
        self.calls.append((local_path, object_key))
        return f"https://example.oss/{object_key}"


def _items():
    return [
        {
            "offer_id": "TEST-001",
            "name": "Коврик придверный 50x80 см",
            "brand": "Нет бренда",
            "color": "белый",
            "description": "<b>Тест</b> коврик.",
            "mv_length_cm": 50,
            "mv_width_cm": 80,
            "mv_height_cm": 1,
            "mv_weight_kg": 1.2,
            "oss_images_json": [
                {"url": "https://oss.example/a/01.jpg"},
                {"url": "https://oss.example/a/02.jpg"},
            ],
        },
        {
            "offer_id": "TEST-002",
            "name": "Коврик придверный 50x80 синий",
            "color": "синий",
            "oss_images_json": [{"url": "https://oss.example/b/01.jpg"}],
        },
    ]


class TestBuildTemplateRows:
    def test_row_count_and_cells(self):
        rows = build_template_rows(_items())
        assert len(rows) == 2

        r0 = rows[0]
        assert r0[COLS["sku_code"]] == "TEST-001"
        assert r0[COLS["vendor_code"]] == "TEST-001"
        assert r0[COLS["color"]] == "белый"
        assert r0[COLS["fbs"]] == "Да"
        assert r0[COLS["barcode"]] == "Сгенерировать"
        # main photo = first url; html stripped from description
        assert r0[COLS["main_photo"]] == "https://oss.example/a/01.jpg"
        assert r0[72] == "https://oss.example/a/02.jpg"  # first extra photo
        assert "<b>" not in r0[COLS["description"]]
        assert r0[COLS["weight"]] == 1.2

        assert rows[1][COLS["sku_code"]] == "TEST-002"
        assert rows[1][COLS["brand"]] == "Нет бренда"  # default


class TestWriteTemplateXlsx:
    def test_writes_file(self, tmp_path):
        if not TEMPLATE.is_file():
            import pytest
            pytest.skip("blank M.Video template missing")
        out = tmp_path / "filled.xlsx"
        rows = build_template_rows(_items())
        write_template_xlsx(rows, str(TEMPLATE), str(out))
        assert out.is_file()

        wb = openpyxl.load_workbook(out)
        ws = wb["Шаблон для загрузки товаров"]
        r = DATA_START_ROW
        assert ws.cell(row=r, column=COLS["sku_code"]).value == "TEST-001"
        assert ws.cell(row=r + 1, column=COLS["sku_code"]).value == "TEST-002"


class TestImagePipeline:
    def test_process_local_images_3x4(self, tmp_path):
        src = tmp_path / "src"
        src.mkdir()
        (Image.new("RGB", (500, 500), (200, 0, 0))).save(src / "p1.jpg")
        out_dir = tmp_path / "out"
        paths = process_local_images([str(src / "p1.jpg")], str(out_dir), offer_id="TEST-001")
        assert len(paths) == 1
        im = Image.open(paths[0])
        assert (im.width, im.height) == (900, 1200)

    def test_upload_glue_manifest(self):
        up = FakeUploader()
        manifest = upload_processed_to_oss(
            up, ["/tmp/01.jpg", "/tmp/02.jpg"], "mvideo/TEST-001"
        )
        assert [m["url"] for m in manifest] == [
            "https://example.oss/mvideo/TEST-001/01.jpg",
            "https://example.oss/mvideo/TEST-001/02.jpg",
        ]
        assert up.calls[0] == ("/tmp/01.jpg", "mvideo/TEST-001/01.jpg")
