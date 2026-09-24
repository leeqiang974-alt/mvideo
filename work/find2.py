import openpyxl, sys
sys.stdout.reconfigure(encoding="utf-8")
for f in ["template_container","template_kaleidoscope"]:
    wb=openpyxl.load_workbook(rf"E:\mvideo\MvideoERP\work\downloads\{f}.xlsx")
    sn="Шаблон для загрузки товаров" if "Шаблон для загрузки товаров" in wb.sheetnames else wb.sheetnames[1]
    ws=wb[sn]
    main=[c for c in range(1,ws.max_column+1) if "Главная фотография" in str(ws.cell(4,c).value)]
    print(f,"cols",ws.max_column,"main",main)
