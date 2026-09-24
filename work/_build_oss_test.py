# -*- coding: utf-8 -*-
import sys, io, json, re, shutil
from pathlib import Path
import httpx, oss2
from PIL import Image
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
ROOT=Path(r"E:\mvideo\MvideoERP\work")
PILOT=json.loads((ROOT/"_test_ru.json").read_text(encoding="utf-8"))
TPL=ROOT/"downloads"/"template_brosh.xlsx"
OUT=ROOT/"downloads"/"_test_oss.xlsx"
lines=Path(r"D:\Desktop\api\阿里云的key和secret.txt").read_text(encoding="utf-8").splitlines()
bucket=oss2.Bucket(oss2.Auth(lines[1].strip(),lines[4].strip()),"https://oss-cn-shanghai.aliyuncs.com","ozonshanghai",connect_timeout=20)
def pad(b):
    im=Image.open(io.BytesIO(b)).convert("RGB"); w,h=im.size; t=3/4; cur=w/h
    if cur<t:
        nw=int(h*t); c=Image.new("RGB",(nw,h),"white"); c.paste(im,((nw-w)//2,0)); im=c
    elif cur>t:
        nh=int(w/t); c=Image.new("RGB",(w,nh),"white"); c.paste(im,(0,(nh-h)//2)); im=c
    return im.resize((900,1200),Image.LANCZOS)
def host(u,offer,k):
    key=f"mvideo-test/{offer}_{k}.jpg"
    try:
        b=httpx.get(u,timeout=30,follow_redirects=True).content
        im=pad(b); buf=io.BytesIO(); im.save(buf,"JPEG",quality=92)
        bucket.put_object(key,buf.getvalue())
        return f"https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/{key}"
    except Exception as e:
        print("IMG FAIL",offer,k,repr(e)[:80]); return None
def strip_html(s):
    s=re.sub(r"<br\s*/?>","\n",s or ""); s=re.sub(r"<[^>]+>","",s)
    return s.replace("&quot;",'"').replace("&amp;","&").strip()
shutil.copy2(TPL,OUT); wb=openpyxl.load_workbook(OUT); ws=wb["Шаблон для загрузки товаров"]
for n,p in enumerate(PILOT):
    r=5+n; offer=p["offer_id"]
    imgs=[p["primary_image"]]+[u for u in p.get("images",[]) if u!=p["primary_image"]]
    urls=[host(u,offer,k) for k,u in enumerate(imgs[:5])]
    urls=[u for u in urls if u]
    desc=strip_html(p["description_html"])[:1500]
    ws.cell(row=r,column=1,value="Сгенерировать"); ws.cell(row=r,column=2,value=p["name"][:249])
    ws.cell(row=r,column=3,value="Нет бренда"); ws.cell(row=r,column=4,value="Нет")
    ws.cell(row=r,column=5,value="Да"); ws.cell(row=r,column=6,value="Нет")
    ws.cell(row=r,column=7,value=offer[:80]); ws.cell(row=r,column=8,value="Разноцветный")
    ws.cell(row=r,column=9,value=offer); ws.cell(row=r,column=10,value=offer)
    ws.cell(row=r,column=11,value="Китай"); ws.cell(row=r,column=12,value="Нет")
    ws.cell(row=r,column=13,value=0); ws.cell(row=r,column=14,value=10.0)
    ws.cell(row=r,column=15,value=8.0); ws.cell(row=r,column=16,value=2.0); ws.cell(row=r,column=17,value=0.05)
    ws.cell(row=r,column=43,value="Нет"); ws.cell(row=r,column=44,value="Нет"); ws.cell(row=r,column=45,value="0")
    ws.cell(row=r,column=47,value=desc); ws.cell(row=r,column=52,value="Металл")
    ws.cell(row=r,column=56,value="Разноцветный"); ws.cell(row=r,column=70,value="Брошь-булавка")
    ws.cell(row=r,column=72,value="Унисекс")
    for k,u in enumerate(urls[:15]): ws.cell(row=r,column=75+k,value=u)
    print("row",r,offer,"photos",len(urls))
wb.save(OUT); print("SAVED",OUT)
