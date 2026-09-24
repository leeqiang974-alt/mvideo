import openpyxl, glob, sys
sys.stdout.reconfigure(encoding="utf-8")
for f in ["template_carpet","template_hat","template_suitcase","template_scrunchie"]:
    p=rf"E:\mvideo\MvideoERP\work\downloads\{f}.xlsx"
    wb=openpyxl.load_workbook(p)
    sn="Шаблон для загрузки товаров" if "Шаблон для загрузки товаров" in wb.sheetnames else wb.sheetnames[1]
    ws=wb[sn]
    main=[c for c in range(1,ws.max_column+1) if "Главная фотография" in str(ws.cell(4,c).value)]
    req=[c for c in range(1,20) if ws.cell(4,c).value and "*" in str(ws.cell(4,c).value)]
    print(f, "cols",ws.max_column,"main_photo",main)
