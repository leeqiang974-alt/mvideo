import openpyxl,sys
sys.stdout.reconfigure(encoding="utf-8")
wb=openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\downloads\_test_ru_clean.xlsx",read_only=True,data_only=True)
ws=wb["Шаблон для загрузки товаров"]
for r in range(5,ws.max_row+1):
    o=ws.cell(row=r,column=7).value
    if o: print(o, "|", ws.cell(row=r,column=75).value)
