# -*- coding: utf-8 -*-
"""Generic brooch batch builder: input json -> images OSS -> filled xlsx."""
import json, re, os, io, shutil, sys
from pathlib import Path
import httpx, oss2
from PIL import Image
import openpyxl

IN = Path(sys.argv[1]) if len(sys.argv)>1 else Path(r"E:\mvideo\MvideoERP\work\brooch_batch.json")
OUT = Path(sys.argv[2]) if len(sys.argv)>2 else Path(r"E:\mvideo\MvideoERP\work\downloads\mv_upload_brooch_b2.xlsx")
TPL = Path(r"E:\mvideo\MvideoERP\work\downloads\template_brosh.xlsx")
products = json.loads(IN.read_text(encoding="utf-8"))

lines = Path(r"D:\Desktop\api\阿里云的key和secret.txt").read_text(encoding="utf-8").splitlines()
auth = oss2.Auth(lines[1].strip(), lines[4].strip())
bucket = oss2.Bucket(auth, "https://oss-cn-shanghai.aliyuncs.com", "ozonshanghai", connect_timeout=20)

def strip_html(s):
    s = re.sub(r"<br\s*/?>","\n",s or ""); s=re.sub(r"<[^>]>","",s); return s.replace("&quot;",'"').replace("&amp;","&").strip()
def pad_3x4(b):
    im=Image.open(io.BytesIO(b)).convert("RGB"); w,h=im.size; t=3/4; c=w/h
    if c<t: nw=int(h*t); cv=Image.new("RGB",(nw,h),"white"); cv.paste(im,((nw-w)//2,0)); im=cv
    elif c>t: nh=int(w/t); cv=Image.new("RGB",(w,nh),"white"); cv.paste(im,(0,(nh-h)//2)); im=cv
    im=im.resize((900,1200),Image.LANCZOS); buf=io.BytesIO(); im.save(buf,"JPEG",quality=92); return buf.getvalue()
def up(url,offer,k):
    r=httpx.get(url,timeout=30,follow_redirects=True); r.raise_for_status()
    obj=f"mvideo/brooch/{offer}/{k:02d}.jpg"
    bucket.put_object(obj,pad_3x4(r.content),headers={"Content-Type":"image/jpeg"})
    return f"https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/{obj}"

shutil.copy2(TPL,OUT); wb=openpyxl.load_workbook(OUT); ws=wb["Шаблон для загрузки товаров"]
for n,p in enumerate(products):
    r=5+n; offer=p["offer_id"]
    imgs=[p["primary_image"]]+[u for u in p.get("images",[]) if u!=p["primary_image"]]
    urls=[]
    for k,u in enumerate(imgs[:5]):
        try: urls.append(up(u,offer,k))
        except Exception as e: print("IMGFAIL",offer,k,repr(e)[:80])
    ws.cell(row=r,column=1,value="Сгенерировать")
    ws.cell(row=r,column=2,value=p["name"][:249])
    ws.cell(row=r,column=3,value="Нет бренда")
    ws.cell(row=r,column=4,value="Нет"); ws.cell(row=r,column=5,value="Да"); ws.cell(row=r,column=6,value="Нет")
    ws.cell(row=r,column=7,value=offer[:80]); ws.cell(row=r,column=8,value="Разноцветный")
    ws.cell(row=r,column=9,value=offer); ws.cell(row=r,column=10,value=offer)
    ws.cell(row=r,column=11,value="Китай"); ws.cell(row=r,column=12,value="Нет"); ws.cell(row=r,column=13,value=0)
    ws.cell(row=r,column=14,value=10.0); ws.cell(row=r,column=15,value=8.0)
    ws.cell(row=r,column=16,value=2.0); ws.cell(row=r,column=17,value=0.05)
    ws.cell(row=r,column=43,value="Нет"); ws.cell(row=r,column=44,value="Нет"); ws.cell(row=r,column=45,value="0")
    ws.cell(row=r,column=47,value=strip_html(p["description_html"])[:2000])
    ws.cell(row=r,column=52,value="Металл"); ws.cell(row=r,column=56,value="Разноцветный")
    ws.cell(row=r,column=70,value="Брошь-булавка"); ws.cell(row=r,column=72,value="Унисекс")
    for k,u in enumerate(urls[:15]): ws.cell(row=r,column=75+k,value=u)
    print("row",r,offer,"photos",len(urls))
wb.save(OUT); print("SAVED",OUT)
