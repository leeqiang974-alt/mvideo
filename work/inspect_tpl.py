import openpyxl
p = r"E:\mvideo\MvideoERP\work\downloads\template_brosh.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
print("sheets:", wb.sheetnames)
ws = wb["Шаблон для загрузки товаров"]
print("dims:", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
# header rows: row1 = field code, row2 = russian name, row3 = note
for r in range(1, 6):
    print(f"--- ROW {r} ---")
    for c in range(1, ws.max_column+1):
        v = ws.cell(row=r, column=c).value
        if v not in (None, ""):
            print(f"  c{c}: {str(v)[:75]}")
