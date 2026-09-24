# -*- coding: utf-8 -*-
"""Build the brooch pilot Excel: images -> 3:4 OSS -> fill template."""
import json, re, os, io, shutil, hashlib
from pathlib import Path
import httpx, oss2
from PIL import Image
import openpyxl

import sys
ROOT = Path(r"E:\mvideo\MvideoERP\work")
_in = sys.argv[1] if len(sys.argv)>1 else "brooch_pilot.json"
_out = sys.argv[2] if len(sys.argv)>2 else "mv_upload_brooch_pilot.xlsx"
PILOT = json.loads((ROOT / _in).read_text(encoding="utf-8"))
TPL = ROOT / "downloads" / "template_brosh.xlsx"
OUT = ROOT / "downloads" / _out
IMG_DIR = ROOT / "brooch_imgs"
IMG_DIR.mkdir(exist_ok=True)

# OSS creds
lines = Path(r"D:\Desktop\api\阿里云的key和secret.txt").read_text(encoding="utf-8").splitlines()
AKID = lines[1].strip()
SECRET = lines[4].strip()
ENDPOINT = "oss-cn-shanghai.aliyuncs.com"
BUCKET = "ozonshanghai"

auth = oss2.Auth(AKID, SECRET)
bucket = oss2.Bucket(auth, f"https://{ENDPOINT}", BUCKET, connect_timeout=20)


def strip_html(s):
    s = re.sub(r"<br\s*/?>", "\n", s or "")
    s = re.sub(r"<[^>]+>", "", s)
    return s.replace("&quot;", '"').replace("&amp;", "&").strip()


def pad_3x4(src_bytes):
    im = Image.open(io.BytesIO(src_bytes)).convert("RGB")
    w, h = im.size
    target = 3 / 4
    cur = w / h
    if cur < target:
        nw = int(h * target); canvas = Image.new("RGB", (nw, h), "white"); canvas.paste(im, ((nw - w)//2, 0)); im = canvas
    elif cur > target:
        nh = int(w / target); canvas = Image.new("RGB", (w, nh), "white"); canvas.paste(im, (0, (nh - h)//2)); im = canvas
    im = im.resize((900, 1200), Image.LANCZOS)
    buf = io.BytesIO(); im.save(buf, "JPEG", quality=92); return buf.getvalue()


_CACHE = {}
def upload(url, offer, key):
    # v2 (2026-09-22): download -> pad white 900x1200 (3:4) -> host Shanghai OSS -> public URL.
    # Direct Ozon URLs fail: 1:1 aspect rejected and RU backend cannot fetch CN .cn reliably.
    if not url:
        return None
    if url in _CACHE:
        return _CACHE[url]
    last = None
    for _try in range(3):
        try:
            b = httpx.get(url, timeout=30, follow_redirects=True).content
            data = pad_3x4(b)
            name = hashlib.md5(data).hexdigest()
            okey = f"mvideo/img/{name}.jpg"
            bucket.put_object(okey, data)
            out = f"https://{BUCKET}.{ENDPOINT}/{okey}"
            _CACHE[url] = out
            return out
        except Exception as e:
            last = e
    raise last


# copy template and fill
shutil.copy2(TPL, OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb["Шаблон для загрузки товаров"]

for n, p in enumerate(PILOT):
    r = 5 + n
    offer = p["offer_id"]
    imgs = [p["primary_image"]] + [u for u in p.get("images", []) if u != p["primary_image"]]
    # upload main + up to 4 more
    urls = []
    for k, u in enumerate(imgs[:5]):
        try:
            urls.append(upload(u, offer, f"{k:02d}"))
        except Exception as e:
            print("IMG FAIL", offer, k, repr(e)[:80])
    desc = strip_html(p["description_html"])[:1500]
    ws.cell(row=r, column=1, value="Сгенерировать")
    ws.cell(row=r, column=2, value=p["name"][:249])
    ws.cell(row=r, column=3, value="Нет бренда")
    ws.cell(row=r, column=4, value="Нет")
    ws.cell(row=r, column=5, value="Да")
    ws.cell(row=r, column=6, value="Нет")
    ws.cell(row=r, column=7, value=offer[:80])
    ws.cell(row=r, column=8, value="Разноцветный")
    ws.cell(row=r, column=9, value=offer)
    ws.cell(row=r, column=10, value=offer)
    ws.cell(row=r, column=11, value="Китай")
    ws.cell(row=r, column=12, value="Нет")
    ws.cell(row=r, column=13, value=0)
    ws.cell(row=r, column=14, value=10.0)
    ws.cell(row=r, column=15, value=8.0)
    ws.cell(row=r, column=16, value=2.0)
    ws.cell(row=r, column=17, value=0.05)
    ws.cell(row=r, column=43, value="Нет")
    ws.cell(row=r, column=44, value="Нет")
    ws.cell(row=r, column=45, value="0")
    ws.cell(row=r, column=47, value=desc)
    ws.cell(row=r, column=52, value="Металл")
    ws.cell(row=r, column=56, value="Разноцветный")
    ws.cell(row=r, column=70, value="Брошь-булавка")
    ws.cell(row=r, column=72, value="Унисекс")
    # photos
    for k, u in enumerate(urls[:15]):
        ws.cell(row=r, column=75 + k, value=u)
    print("filled row", r, offer, "photos", len(urls))

wb.save(OUT)
print("SAVED", OUT)

