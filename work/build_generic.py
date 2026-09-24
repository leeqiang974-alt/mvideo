# -*- coding: utf-8 -*-
import json, sys, io, shutil
from pathlib import Path
import httpx, oss2
from PIL import Image
import openpyxl
# args: input.json output.xlsx template.xlsx MAINCOL
inp,out,tpl,maincol=sys.argv[1],sys.argv[2],sys.argv[3],int(sys.argv[4])
products=json.loads(Path(inp).read_text(encoding="utf-8"))
lines=Path(r"D:\Desktop\api\阿里云的key和secret.txt").read_text(encoding="utf-8").splitlines()
auth=oss2.Auth(lines[1].strip(),lines[4].strip())
bucket=oss2.Bucket(auth,"https://oss-cn-shanghai.aliyuncs.com","ozonshanghai",connect_timeout=20)
def pad(b):
    im=Image.open(io.BytesIO(b)).convert("RGB"); w,h=im.size; t=3/4; c=w/h
    if c<t: nw=int(h*t); cv=Image.new("RGB",(nw,h),"white"); cv.paste(im,((nw-w)//2,0)); im=cv
    elif c>t: nh=int(w/t); cv=Image.new("RGB",(w,nh),"white"); cv.paste(im,(0,(nh-h)//2)); im=cv
    im=im.resize((900,1200),Image.LANCZOS); buf=io.BytesIO(); im.save(buf,"JPEG",quality=92); return buf.getvalue()
def up(url,offer,k):
    r=httpx.get(url,timeout=30,follow_redirects=True); r.raise_for_status()
    import hashlib; slug=hashlib.md5(offer.encode()).hexdigest()[:16]
    obj=f"mvideo/{Path(out).stem}/{slug}/{k:02d}.jpg"
    bucket.put_object(obj,pad(r.content),headers={"Content-Type":"image/jpeg"})
    return f"https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/{obj}"
shutil.copy2(tpl,out); wb=openpyxl.load_workbook(out); ws=wb["Шаблон для загрузки товаров"]
for n,p in enumerate(products):
    r=5+n; offer=p["offer_id"]
    urls=[]
    for k,u in enumerate(p["imgs"][:5]):
        try: urls.append(up(u,offer,k))
        except Exception as e: print("IMGFAIL",offer,k,repr(e)[:60])
    ws.cell(r,1,"Сгенерировать"); ws.cell(r,2,p["name"][:249]); ws.cell(r,3,"Нет бренда")
    ws.cell(r,4,"Нет"); ws.cell(r,5,"Да"); ws.cell(r,6,"Нет")
    ws.cell(r,7,offer[:80]); ws.cell(r,8,"Разноцветный"); ws.cell(r,9,offer); ws.cell(r,10,offer)
    ws.cell(r,11,"Китай"); ws.cell(r,12,"Нет"); ws.cell(r,13,0)
    ws.cell(r,14,10.0); ws.cell(r,15,8.0); ws.cell(r,16,2.0); ws.cell(r,17,0.05)
    ws.cell(r,43,"Нет"); ws.cell(r,44,"Нет"); ws.cell(r,45,"0")
    for k,u in enumerate(urls): ws.cell(r,maincol+k,u)
    print("row",r,offer,"photos",len(urls))
wb.save(out); print("SAVED",out)

