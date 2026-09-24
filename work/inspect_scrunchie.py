import openpyxl
wb=openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\downloads\template_scrunchie.xlsx")
print("sheets",wb.sheetnames)
for sn in wb.sheetnames:
    ws=wb[sn]
    print(sn,ws.max_row,ws.max_column)
ws=wb["Шаблон для загрузки товаров"] if "Шаблон для загрузки товаров" in wb.sheetnames else wb[wb.sheetnames[1]]
print("DATA SHEET:",ws.title)
for c in range(1,ws.max_column+1):
    print(c,"|",str(ws.cell(4,c).value)[:50])
