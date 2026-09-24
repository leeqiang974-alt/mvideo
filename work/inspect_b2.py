# -*- coding: utf-8 -*-
import openpyxl
pilot=openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\downloads\mv_upload_brooch_pilot.xlsx")
b2=openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\downloads\mv_upload_brooch_b2.xlsx")
ps=pilot.active; bs=b2.active
print("pilot dims", ps.max_row, ps.max_column, "| b2 dims", bs.max_row, bs.max_column)
# header row is row 4
ph=[ps.cell(4,c).value for c in range(1,ps.max_column+1)]
bh=[bs.cell(4,c).value for c in range(1,bs.max_column+1)]
print("cols equal:", ph==bh)
if ph!=bh:
    for i,(a,c) in enumerate(zip(ph,bh),1):
        if a!=c: print("  DIFF col",i,repr(a),"vs",repr(c))
# count rows and image URLs (col75 main image)
img_urls=[]
for r in range(5,bs.max_row+1):
    url=bs.cell(r,75).value
    offer=bs.cell(r,9).value
    img_urls.append((r,offer,url))
print("b2 data rows:", len(img_urls))
print("sample:", img_urls[0])
import json
json.dump([{"row":r,"offer":o,"url":u} for r,o,u in img_urls],
          open(r"E:\mvideo\MvideoERP\work\b2_urls.json","w",encoding="utf-8"),ensure_ascii=False,indent=1)
