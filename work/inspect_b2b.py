import openpyxl, json
wb=openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\downloads\mv_upload_brooch_b2.xlsx")
ws=wb["Шаблон для загрузки товаров"]
hdr={c: ws.cell(4,c).value for c in range(1,95)}
# required cols
req={1:"c1",2:"name",3:"brand",7:"model",9:"offer",14:"L",15:"W",16:"H",17:"wt",75:"img"}
rows=[]
for r in range(5,ws.max_row+1):
    rec={}
    for c in range(1,95):
        v=ws.cell(r,c).value
        if v not in (None,""): rec[c]=str(v)[:40]
    rows.append((r,rec))
print("data rows:",len(rows))
# check col75 (main image) presence and any None required
for r,rec in rows:
    missing=[c for c in [1,2,3,7,9,14,15,16,17,75] if c not in rec]
    img=rec.get(75,"")
    print(r,"offer=",rec.get(9),"img_ok=",img.startswith("http"),"missing=",missing)
json.dump([{"row":r,"offer":rec.get(9),"img":rec.get(75)} for r,rec in rows],
          open(r"E:\mvideo\MvideoERP\work\b2_urls.json","w",encoding="utf-8"),ensure_ascii=False,indent=1)
