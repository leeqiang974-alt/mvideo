import openpyxl
wb=openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\downloads\mv_upload_brooch_b2.xlsx")
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws=wb[sn]
    print(sn, ws.max_row, ws.max_column)
