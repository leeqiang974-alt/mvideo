# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\out_ozon2mv.xlsx", read_only=True)
ws = wb[wb.sheetnames[1]]
rows = list(ws.iter_rows(values_only=True))
header = rows[3]
data = rows[4]
targets = [0, 1, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 19, 42, 43, 44, 45, 57, 70, 71]
for i in targets:
    print(f"[{i}] {header[i]} = {data[i]!r}")
