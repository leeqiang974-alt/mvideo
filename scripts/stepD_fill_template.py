# -*- coding: utf-8 -*-
"""Step D: fill M.Video Excel template with 8 SKU00259 variants.

Reads:
  - work/mvideo_template.xlsx (downloaded template, do NOT restructure)
  - work/oss_urls_manifest.json (processed+uploaded image URLs)
  - sku00259_full.json (source attributes)

Writes:
  - work/mvideo_template_filled.xlsx
Data rows start at row 5 (rows 1-4 are header blocks).
"""
import json
import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(r"E:\mvideo\MvideoERP")
TEMPLATE = ROOT / "work" / "mvideo_template.xlsx"
OUT = ROOT / "work" / "mvideo_template_filled.xlsx"
OSS_MANIFEST = ROOT / "work" / "oss_urls_manifest.json"
SRC_JSON = ROOT / "sku00259_full.json"

# Color mapping per offer_id (Russian, matching M.Video color list or close)
COLOR_MAP = {
    "SKU00259": "белый",
    "SKU00259-Светло-жёлтый": "светло-желтый",
    "SKU00259-Светло-жёлтый-hi": "светло-желтый",
    "SKU00259-Синий-hello": "синий",
    "SKU00259-Белый-hi": "белый",
    "SKU00259-оранжевый-hello": "оранжевый",
    "SKU00259-Зелёный-hello": "зеленый",
    "SKU00259-Жёлтый смайлик": "желтый",
}

# Short description (strip HTML, trim to 1500 chars)
def clean_html(t: str) -> str:
    t = re.sub(r"<br\s*/?>", "\n", t or "")
    t = re.sub(r"<[^>]+>", "", t)
    t = re.sub(r"&[a-z]+;", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t[:1500]


def main():
    src = json.loads(SRC_JSON.read_text(encoding="utf-8"))
    oss = json.loads(OSS_MANIFEST.read_text(encoding="utf-8"))

    # Build lookup: offer_id -> description
    desc_by_offer = {}
    for r in src["attributes"]["result"]:
        offer = r["offer_id"]
        for a in r.get("attributes", []):
            if a.get("id") == 4191:
                vals = a.get("values", [])
                if vals:
                    desc_by_offer[offer] = clean_html(vals[0].get("value", ""))
                break

    # Build lookup: offer_id -> weight
    weight_by_offer = {}
    for r in src["attributes"]["result"]:
        weight_by_offer[r["offer_id"]] = round((r.get("weight", 1000) or 1000) / 1000.0, 3)

    # Copy template
    shutil.copy2(TEMPLATE, OUT)
    wb = openpyxl.load_workbook(OUT)
    ws = wb["Шаблон для загрузки товаров"]

    # Column indices (1-based)
    COLS = {
        "barcode": 1,       # A
        "name": 2,          # B
        "brand": 3,         # C
        "dbs": 4,           # D
        "fbs": 5,           # E
        "fbm": 6,           # F
        "model": 7,         # G
        "color": 8,         # H
        "sku_code": 9,     # I
        "vendor_code": 10,  # J
        "country": 11,     # K
        "warranty": 12,    # L
        "shelf_life": 13,  # M
        "length": 14,      # N
        "width": 15,       # O
        "height": 16,      # P
        "weight": 17,      # Q
        "vertical": 43,    # AQ
        "fragile": 44,     # AR
        "nds": 45,         # AS
        "material": 55,    # BC
        "material_base": 57,  # BE
        "description": 58,   # BF
        "main_photo": 71,    # BS
    }
    # Photos 1..14 -> cols 72..85 (BT..CG)
    PHOTO_COLS = list(range(72, 86))

    row = 5
    for variant in oss:
        offer = variant["offer_id"]
        color = COLOR_MAP.get(offer, "белый")
        images = variant["images"]

        # Write cells
        ws.cell(row=row, column=COLS["barcode"], value="Сгенерировать")
        ws.cell(row=row, column=COLS["name"],
                value=f"Коврик придверный 50x80 см, {color}")
        ws.cell(row=row, column=COLS["brand"], value="Нет бренда")
        ws.cell(row=row, column=COLS["dbs"], value="Нет")
        ws.cell(row=row, column=COLS["fbs"], value="Да")
        ws.cell(row=row, column=COLS["fbm"], value="Нет")
        ws.cell(row=row, column=COLS["model"], value=f"Коврик придверный 50x80 {color}")
        ws.cell(row=row, column=COLS["color"], value=color)
        ws.cell(row=row, column=COLS["sku_code"], value=offer)
        ws.cell(row=row, column=COLS["vendor_code"], value=offer)
        ws.cell(row=row, column=COLS["country"], value="Китай")
        ws.cell(row=row, column=COLS["warranty"], value="Нет")
        ws.cell(row=row, column=COLS["shelf_life"], value=0)
        ws.cell(row=row, column=COLS["length"], value=50.0)
        ws.cell(row=row, column=COLS["width"], value=15.0)
        ws.cell(row=row, column=COLS["height"], value=15.0)
        ws.cell(row=row, column=COLS["weight"], value=weight_by_offer.get(offer, 1.0))
        ws.cell(row=row, column=COLS["vertical"], value="Нет")
        ws.cell(row=row, column=COLS["fragile"], value="Нет")
        ws.cell(row=row, column=COLS["nds"], value="0")
        ws.cell(row=row, column=COLS["material"], value="ПВХ")
        ws.cell(row=row, column=COLS["material_base"], value="резина")
        ws.cell(row=row, column=COLS["description"],
                value=desc_by_offer.get(offer, "Придверный коврик 50x80 см."))

        # Main photo = first image
        if images:
            ws.cell(row=row, column=COLS["main_photo"], value=images[0]["url"])
            # Remaining images -> photo 1..14
            for i, img in enumerate(images[1:15]):
                col_idx = PHOTO_COLS[i]
                ws.cell(row=row, column=col_idx, value=img["url"])

        print(f"  row {row}: {offer} ({color})  photos={len(images)}")
        row += 1

    wb.save(OUT)
    print(f"\nDONE. filled {row-5} rows -> {OUT}")


if __name__ == "__main__":
    main()
