# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\app1_commission.xlsx", read_only=True)
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n=== sheet: {sn} dims={ws.max_row}x{ws.max_column} ===")
    count = 0
    for row in ws.iter_rows(values_only=True):
        vals = [str(v) for v in row if v is not None]
        if vals:
            print(" | ".join(vals)[:200])
            count += 1
        if count >= 25:
            print("... (truncated)")
            break
