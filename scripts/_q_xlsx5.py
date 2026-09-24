# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\out_ozon2mv.xlsx", read_only=True)
ws = wb[wb.sheetnames[1]]
rows = list(ws.iter_rows(values_only=True))
header = rows[3]
data = None
for r in rows[4:]:
    if r[0]:
        data = r
        break
print("sheet:", wb.sheetnames[1])
print("修改时间检查用文件系统")
for i in [0, 1, 2, 9, 13, 14, 15, 16, 58, 70, 71]:
    print(f"[{i}] {header[i]} = {data[i]!r}")
