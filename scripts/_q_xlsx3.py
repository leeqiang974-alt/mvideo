# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\out_ozon2mv.xlsx", read_only=True)
ws = wb[wb.sheetnames[1]]
rows = list(ws.iter_rows(values_only=True))
header = rows[3]
# find data row (first row with barcode value)
data = None
for r in rows[4:]:
    if r[0]:
        data = r
        break
print("data row found:", data is not None)
targets = [0, 1, 9, 13, 14, 15, 16, 45, 57, 70]
for i in targets:
    print(f"[{i}] {header[i]} = {data[i]!r}")
