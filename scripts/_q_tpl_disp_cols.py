# -*- coding: utf-8 -*-
"""输出切带机模板第1-4行的完整列头（95列）。"""
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\template_dispenser.xlsx", read_only=False)
ws = wb["Шаблон для загрузки товаров"]
print("dims:", ws.max_row, "x", ws.max_column)
for row_idx in [1, 2, 3, 4]:
    print(f"\n--- R{row_idx} ---")
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=row_idx, column=c).value
        if v is not None:
            print(f"  C{c}: {str(v)[:60]}")
