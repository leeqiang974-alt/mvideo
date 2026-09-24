# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\out_ozon2mv.xlsx", read_only=True)
ws = wb[wb.sheetnames[1]]
rows = list(ws.iter_rows(values_only=True))
header = rows[3]
data = None
for r in rows[4:]:
    if r[0]:
        data = r
        break
targets = [0, 1, 13, 14, 15, 16, 57, 58, 70]
for i in targets:
    v = data[i]
    print(f"[{i}] {header[i]} = {str(v)[:120]!r}")
