# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\out_ozon2mv.xlsx", read_only=True)
ws = wb[wb.sheetnames[1]]
rows = list(ws.iter_rows(values_only=True))
print("sheet:", wb.sheetnames[1], "total rows:", len(rows))
# header row 4, data row 5
header = rows[3]
data = rows[4]
print("=== header ===")
for i, h in enumerate(header):
    if h is not None:
        print(f"[{i}] {h}")
print("=== data row (row5) ===")
for i, h in enumerate(header):
    v = data[i]
    if h is not None:
        print(f"[{i}] {h} = {v}")
