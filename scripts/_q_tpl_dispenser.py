# -*- coding: utf-8 -*-
"""解析切带机模板：sheet 结构、头部说明（groupCode/INF）、列名。"""
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\template_dispenser.xlsx", read_only=False)
print("SHEETS:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n=== sheet: {sn} dims={ws.max_row}x{ws.max_column} ===")
    # 打印前 12 行（通常含说明/列头）
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i > 12: break
        vals = [str(v)[:40] if v is not None else "" for v in row]
        nonempty = [v for v in vals if v]
        if nonempty:
            print(f"R{i}:", " | ".join(nonempty)[:250])
