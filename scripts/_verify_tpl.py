# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
wb = openpyxl.load_workbook(r"C:\OzonERP\verify_tpl.xlsx", read_only=True)
for ws in wb.worksheets:
    print("SHEET:", ws.title, ws.max_row, "x", ws.max_column)
ws = wb.active
rows = list(ws.iter_rows(min_row=1, max_row=10, values_only=True))
print("HEADER:", [str(c)[:18] if c else "" for c in rows[0]][:22])
for r in rows[1:]:
    print("ROW:", [str(c)[:20] if c is not None else "." for c in r[:14]])
