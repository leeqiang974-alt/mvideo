# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"E:\mvideo\MvideoERP\work\out_ozon2mv.xlsx", read_only=True)
ws = wb[wb.sheetnames[1]]
rows = list(ws.iter_rows(values_only=True))
header = rows[3]
data = None
for r in rows[4:]:
    if r[0]:
        data = r
        break
# dump ALL columns with non-empty data value
print("=== 数据行所有非空列 ===")
for i in range(len(header)):
    if header[i] is not None and data[i] not in (None, ""):
        print(f"[{i}] {header[i]} = {str(data[i])[:80]!r}")
