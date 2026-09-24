"""Single-SKU Ozon -> M.Video dry-run workflow.

This module is intentionally separate from the bulk migration state machine.
The current release builds one Excel-template row and never performs an
upload; source Ozon RUB prices are retained as evidence only.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Mapping
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import CategoryMapping, SingleSkuJob
from .brand_sanitizer import NO_BRAND, sanitize_listing
from .pricing import (
    PricingError,
    calculate_price,
    pricing_json_safe,
    to_decimal,
)

PROJECT_ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_WORKSHEET = "Шаблон для загрузки товаров"
DISPENSER_OZON_CATEGORY_ID = 17029021
DISPENSER_MV_GROUP_ID = "604171101"
DISPENSER_MV_INFOMODEL_ID = "INF-307590"
DISPENSER_TEMPLATE_PATH = "work/template_dispenser.xlsx"


@dataclass(frozen=True)
class CategoryPublishRule:
    """Category-specific publishing requirements and Excel template."""

    ozon_category_id: int
    mv_group_id: str
    mv_infomodel_id: str
    template_file_path: str
    worksheet_name: str = TEMPLATE_WORKSHEET
    commission_category: str = "home"
    tn_ved_required: bool = False
    required_certificates: tuple[str, ...] = ()
    brand_authorization_required: bool = False
    required_material: bool = True


CATEGORY_PUBLISH_RULES: dict[int, CategoryPublishRule] = {
    DISPENSER_OZON_CATEGORY_ID: CategoryPublishRule(
        ozon_category_id=DISPENSER_OZON_CATEGORY_ID,
        mv_group_id=DISPENSER_MV_GROUP_ID,
        mv_infomodel_id=DISPENSER_MV_INFOMODEL_ID,
        template_file_path=DISPENSER_TEMPLATE_PATH,
        commission_category="home",
        tn_ved_required=False,
        required_certificates=(),
        brand_authorization_required=False,
        required_material=True,
    )
}


class SingleSkuWorkflowError(ValueError):
    """Blockers prevented the single-SKU workflow from continuing."""

    def __init__(self, message: str, blockers: list[str] | None = None) -> None:
        super().__init__(message)
        self.blockers = list(blockers or [])


def _is_mapping(value: Any) -> bool:
    return isinstance(value, Mapping)


def _text(value: Any) -> str:
    return str(value or "").strip()


def _first_present(payload: Mapping[str, Any], keys: tuple[str, ...]) -> Any:
    for key in keys:
        if key in payload:
            return payload[key]
    return None


def _positive_decimal(candidate: Any) -> Decimal | None:
    try:
        value = to_decimal(candidate, "value", allow_none=True)
    except PricingError:
        return None
    if value is None or value <= 0:
        return None
    return value


def _source_price_rub(payload: Mapping[str, Any]) -> Decimal | None:
    for key in (
        "source_price_rub",
        "price_rub",
        "price",
        "current_price",
        "webPrice",
    ):
        candidate = payload.get(key)
        if candidate is None:
            continue
        if _is_mapping(candidate):
            for nested_key in ("price", "value", "currentPrice", "cardPrice"):
                value = _positive_decimal(candidate.get(nested_key))
                if value is not None:
                    return value
        else:
            value = _positive_decimal(candidate)
            if value is not None:
                return value
    return None


def _package_decimal(payload: Mapping[str, Any], keys: tuple[str, ...]) -> Decimal | None:
    package = payload.get("packageInfo")
    if not _is_mapping(package):
        package = payload.get("package_info")
    containers = [package] if _is_mapping(package) else []
    containers.append(payload)

    for container in containers:
        candidate = _first_present(container, keys)
        value = _positive_decimal(candidate)
        if value is not None:
            return value
    return None


def _extract_package(payload: Mapping[str, Any]) -> dict[str, Decimal | None]:
    return {
        "length_mm": _package_decimal(payload, ("lengthMm", "length_mm", "length")),
        "width_mm": _package_decimal(payload, ("widthMm", "width_mm", "width")),
        "height_mm": _package_decimal(payload, ("heightMm", "height_mm", "height")),
        "weight_g": _package_decimal(
            payload, ("weightG", "weight_g", "weightGrams", "weight")
        ),
    }


def _image_url(candidate: Any) -> str:
    if _is_mapping(candidate):
        for key in ("url", "src", "original", "image_url", "link"):
            url = _text(candidate.get(key))
            if url:
                return url
        return ""
    return _text(candidate)


def _extract_images(payload: Mapping[str, Any]) -> list[str]:
    candidates: list[Any] = []
    for key in ("images", "images_json", "webGallery", "gallery"):
        value = payload.get(key)
        if isinstance(value, list):
            candidates.extend(value)

    invalid_terms = {"none", "null", "undefined"}
    result: list[str] = []
    for candidate in candidates:
        url = _image_url(candidate)
        if not url.startswith("https://") or url.lower() in invalid_terms:
            continue
        if url not in result:
            result.append(url)
    return result[:15]


def _extract_category(payload: Mapping[str, Any]) -> tuple[int | None, str]:
    category = payload.get("category")
    if not _is_mapping(category):
        category = payload.get("ozon_category")

    category_id: int | None = None
    category_name = ""
    if _is_mapping(category):
        value = _positive_decimal(
            _first_present(category, ("id", "category_id", "ozon_category_id"))
        )
        if value is not None:
            category_id = int(value)
        category_name = _text(
            _first_present(category, ("name", "category_name", "title"))
        )

    if category_id is None:
        value = _positive_decimal(
            _first_present(
                payload,
                ("category_id", "ozon_category_id", "categoryId", "ozonCategoryId"),
            )
        )
        if value is not None:
            category_id = int(value)
    if not category_name:
        category_name = _text(payload.get("category_name"))
    return category_id, category_name


def _yield_attribute_values(label: Any, raw_value: Any):
    if _is_mapping(raw_value):
        for key in ("value", "text", "name"):
            value = _text(raw_value.get(key))
            if value:
                yield label, value
                return
    elif isinstance(raw_value, list):
        for item in raw_value:
            if _is_mapping(item):
                yield from _yield_attribute_values(
                    label, item.get("value", item.get("text", item.get("name")))
                )
            else:
                value = _text(item)
                if value:
                    yield label, value
    else:
        value = _text(raw_value)
        if value:
            yield label, value


def _iter_attributes(payload: Mapping[str, Any]):
    for container_name in (
        "characteristics",
        "webCharacteristics",
        "attributes",
        "aspects",
    ):
        data = payload.get(container_name)
        if _is_mapping(data):
            for label, raw_value in data.items():
                yield from _yield_attribute_values(label, raw_value)
        elif isinstance(data, list):
            for row in data:
                if not _is_mapping(row):
                    continue
                label = row.get("name") or row.get("attributeName") or row.get("title")
                raw_value = row.get("value")
                if raw_value is None:
                    raw_value = row.get("values")
                yield from _yield_attribute_values(label, raw_value)


def _find_attribute(payload: Mapping[str, Any], names: set[str]) -> str:
    for label, value in _iter_attributes(payload):
        if _text(label).lower() in names:
            return value
    return ""


def _extract_material(payload: Mapping[str, Any]) -> str:
    direct = _text(payload.get("material"))
    if direct:
        return direct
    return _find_attribute(
        payload,
        {"material", "материал", "материал корпуса", "материал лезвия"},
    )


def create_job_from_ozon_payload(
    session: Session,
    payload: Mapping[str, Any],
    *,
    job_ref: str | None = None,
) -> SingleSkuJob:
    """Persist source evidence as an independent single-SKU job."""
    if not _is_mapping(payload):
        raise SingleSkuWorkflowError("Ozon payload must be an object", ["来源 payload 必须是对象"])

    package = _extract_package(payload)
    category_id, category_name = _extract_category(payload)
    model = _text(
        _first_present(payload, ("model", "model_name"))
    ) or _find_attribute(payload, {"model", "модель", "article", "артикул"})

    job = SingleSkuJob(
        job_ref=job_ref or f"SSKU-{uuid4().hex[:12].upper()}",
        source_platform="ozon",
        source_url=_text(_first_present(payload, ("url", "source_url", "product_url"))),
        source_product_id=_text(
            _first_present(payload, ("product_id", "source_product_id", "id"))
        ),
        source_sku_id=_text(_first_present(payload, ("sku_id", "source_sku_id"))),
        source_title=_text(_first_present(payload, ("title", "name"))),
        source_description=_text(
            _first_present(payload, ("description", "description_text"))
        ),
        source_brand=_text(_first_present(payload, ("brand", "brand_name"))),
        source_model=model,
        source_price_rub=_source_price_rub(payload),
        source_images_json=_extract_images(payload),
        source_payload_json=dict(payload),
        ozon_category_id=category_id,
        ozon_category_name=category_name,
        length_mm=package["length_mm"],
        width_mm=package["width_mm"],
        height_mm=package["height_mm"],
        weight_g=package["weight_g"],
        status="awaiting_input",
    )
    session.add(job)
    session.commit()
    session.refresh(job)
    return job


def convert_package_dimensions(
    *,
    length_mm: Any,
    width_mm: Any,
    height_mm: Any,
    weight_g: Any,
) -> dict[str, Decimal]:
    """Convert Ozon mm/g units to M.Video cm/kg units exactly."""
    length = to_decimal(length_mm, "length_mm")
    width = to_decimal(width_mm, "width_mm")
    height = to_decimal(height_mm, "height_mm")
    weight = to_decimal(weight_g, "weight_g")
    if any(value <= 0 for value in (length, width, height, weight)):
        raise PricingError("package dimensions and weight must be greater than zero")

    return {
        "length_cm": length / Decimal("10"),
        "width_cm": width / Decimal("10"),
        "height_cm": height / Decimal("10"),
        "weight_kg": weight / Decimal("1000"),
    }


def _block_job(
    session: Session, job: SingleSkuJob, blockers: list[str]
) -> None:
    messages: list[str] = []
    for blocker in blockers:
        message = _text(blocker)
        if message and message not in messages:
            messages.append(message)
    if not messages:
        messages.append("单 SKU 工作流被阻断")
    job.status = "blocked"
    job.last_error = "; ".join(messages)
    session.commit()
    raise SingleSkuWorkflowError(job.last_error, messages)


def _require_confirmed_category(
    session: Session, job: SingleSkuJob, rule: CategoryPublishRule
) -> list[str]:
    if job.ozon_category_id is None:
        return ["缺少 Ozon 类目 ID"]

    stmt = select(CategoryMapping).where(
        CategoryMapping.ozon_category_id == int(job.ozon_category_id)
    )
    mapping = session.execute(stmt).scalar_one_or_none()
    if mapping is None:
        return ["类目映射不存在，必须先建立 confirmed 映射"]
    if _text(mapping.status).lower() != "confirmed":
        return ["类目映射不是 confirmed 状态"]
    if _text(mapping.mv_group_id) != rule.mv_group_id:
        return ["M.Video group_id 与类目发布规则不一致"]
    if _text(mapping.mv_infomodel_id) != rule.mv_infomodel_id:
        return ["M.Video infomodel_id 与类目发布规则不一致"]

    job.mv_group_id = mapping.mv_group_id
    job.mv_group_name = mapping.mv_group_name
    job.mv_infomodel_id = mapping.mv_infomodel_id
    return []


def _normalize_documents(value: Any) -> list[dict[str, str]]:
    if not isinstance(value, (list, tuple)):
        return []

    documents: list[dict[str, str]] = []
    for item in value:
        document: dict[str, str] = {}
        if isinstance(item, str):
            document["type"] = _text(item)
        elif _is_mapping(item):
            document["type"] = _text(
                item.get("cert_type") or item.get("type") or item.get("document_type")
            )
            reg_number = _text(item.get("reg_number") or item.get("number"))
            file_url = _text(item.get("file_url") or item.get("url") or item.get("path"))
            if reg_number:
                document["reg_number"] = reg_number
            if file_url:
                document["file_url"] = file_url
        if document.get("type"):
            documents.append(document)
    return documents


def _compliance_blockers(
    rule: CategoryPublishRule,
    documents: list[dict[str, str]],
    tn_ved: str,
) -> list[str]:
    blockers: list[str] = []
    complete = {
        document["type"]
        for document in documents
        if document.get("reg_number") and document.get("file_url", "").startswith("https://")
    }
    for certificate_type in rule.required_certificates:
        if certificate_type not in complete:
            blockers.append(f"缺少合规证书：{certificate_type}")
    if rule.brand_authorization_required and "brand_authorization" not in complete:
        blockers.append("缺少品牌授权文件")
    if rule.tn_ved_required and not tn_ved:
        blockers.append("缺少 TN VED")
    return blockers


def _resolve_template_path(rule: CategoryPublishRule) -> Path:
    path = Path(rule.template_file_path)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path


def build_dispenser_excel_row(
    job: SingleSkuJob,
    converted: dict[str, Decimal],
    images: list[str],
    material: str,
) -> dict[int, Any]:
    """Build one row for the verified 95-column dispenser template."""
    row: dict[int, Any] = {column: None for column in range(1, 96)}
    row.update(
        {
            1: "Сгенерировать",
            2: job.sanitized_title,
            3: NO_BRAND,
            4: "Нет",
            5: "Да",
            6: "Нет",
            7: job.source_model,
            8: job.color,
            9: job.job_ref,
            10: job.source_model,
            11: "Китай",
            12: "Нет",
            13: 0,
            14: float(converted["length_cm"]),
            15: float(converted["width_cm"]),
            16: float(converted["height_cm"]),
            17: float(converted["weight_kg"]),
            43: "Нет",
            44: "Нет",
            45: 0,
            46: job.tn_ved or "",
            48: float(converted["height_cm"]),
            49: float(converted["width_cm"]),
            50: float(converted["length_cm"]),
            51: job.source_model,
            57: "Общего назначения",
            58: material,
            62: job.sanitized_description,
            64: float(converted["weight_kg"]),
            65: "Общего назначения",
            70: material,
            73: job.color,
            74: 1,
            75: "1 шт",
            76: images[0],
        }
    )
    for index, url in enumerate(images[1:], start=77):
        row[index] = url
    document_urls = [
        doc.get("file_url", "")
        for doc in job.compliance_documents_json
        if doc.get("file_url", "").startswith("https://")
    ]
    for index, url in enumerate(document_urls[:5], start=91):
        row[index] = url
    return row


def render_dry_run_workbook(
    rule: CategoryPublishRule,
    job: SingleSkuJob,
    row: dict[int, Any],
) -> dict[str, Any]:
    """Render row into a copied Excel workbook without uploading it."""
    template_path = _resolve_template_path(rule)
    if not template_path.is_file():
        raise SingleSkuWorkflowError(
            "Excel template is missing", [f"模板文件不存在：{template_path}"]
        )

    workbook = load_workbook(template_path)
    if rule.worksheet_name not in workbook.sheetnames:
        raise SingleSkuWorkflowError(
            "Excel worksheet is missing", [f"模板缺少工作表：{rule.worksheet_name}"]
        )
    worksheet = workbook[rule.worksheet_name]
    if worksheet.max_column < 95:
        raise SingleSkuWorkflowError(
            "Excel template has too few columns", ["模板列数不足 95 列"]
        )

    for column in range(1, 96):
        worksheet.cell(row=5, column=column).value = None
    for column, value in row.items():
        worksheet.cell(row=5, column=column, value=value)

    output_dir = PROJECT_ROOT / "work" / "single_sku"
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_ref = re.sub(r"[^A-Za-z0-9_.-]+", "_", job.job_ref)
    output_path = output_dir / f"dry_run_{job.id}_{safe_ref}.xlsx"
    workbook.save(output_path)

    digest = hashlib.sha256(output_path.read_bytes()).hexdigest()
    non_empty_cells = sorted(
        column for column, value in row.items() if value not in (None, "")
    )
    return {
        "mode": "dry_run",
        "upload_performed": False,
        "upload_ref": "",
        "worksheet_name": rule.worksheet_name,
        "data_row": 5,
        "column_count": 95,
        "template_file_path": rule.template_file_path,
        "output_file_path": str(output_path),
        "output_sha256": digest,
        "non_empty_cells": non_empty_cells,
        "source_price_rub_used_as_cny_cost": False,
        "unit_conversion": {
            "length": "mm / 10 = cm",
            "width": "mm / 10 = cm",
            "height": "mm / 10 = cm",
            "weight": "g / 1000 = kg",
        },
    }


def prepare_dry_run(
    session: Session,
    job: SingleSkuJob,
    *,
    purchase_cost_cny: Any,
    stock: Any,
    color: Any = None,
    target_net_margin: Any = None,
    shipping_channel: Any = None,
    compliance_documents: Any = None,
    tn_ved: Any = "",
) -> SingleSkuJob:
    """Validate confirmed inputs, price, and prepare a dry-run Excel row."""
    blockers: list[str] = []
    rule = CATEGORY_PUBLISH_RULES.get(job.ozon_category_id or -1)
    if rule is None:
        blockers.append("未知类目的发布规则，必须先登记模板和合规要求")

    try:
        cost = to_decimal(purchase_cost_cny, "purchase_cost_cny", allow_none=True)
    except PricingError:
        cost = None
    if cost is None or cost <= 0:
        blockers.append("缺少人工确认且大于 0 的 CNY 采购价")

    try:
        stock_decimal = to_decimal(stock, "stock", allow_none=True)
    except PricingError:
        stock_decimal = None
    if stock_decimal is None or stock_decimal <= 0 or stock_decimal % 1 != 0:
        blockers.append("缺少大于 0 的整数库存")

    for field_name in ("length_mm", "width_mm", "height_mm", "weight_g"):
        value = getattr(job, field_name)
        if value is None or value <= 0:
            blockers.append(f"包装尺重不完整：{field_name}")

    images = []
    for image in job.source_images_json or []:
        url = _text(image)
        if url.startswith("https://") and url not in images:
            images.append(url)
    if not images:
        blockers.append("至少需要一张 HTTPS 商品图片")

    color_value = _text(color) or _text(job.color)
    if not color_value:
        blockers.append("缺少颜色")

    documents = _normalize_documents(compliance_documents)
    tn_ved_value = _text(tn_ved) or _text(job.tn_ved)
    material = _extract_material(job.source_payload_json)

    if rule is not None:
        blockers.extend(_require_confirmed_category(session, job, rule))
        blockers.extend(_compliance_blockers(rule, documents, tn_ved_value))
        if rule.required_material and not material:
            blockers.append("缺少类目必需材质")
        template_path = _resolve_template_path(rule)
        if not template_path.is_file():
            blockers.append(f"模板文件不存在：{template_path}")

    sanitization = sanitize_listing(
        job.source_title,
        job.source_description,
        job.source_brand,
        job.source_model,
    )
    blockers.extend(sanitization.warnings)

    if blockers:
        _block_job(session, job, blockers)

    assert rule is not None
    assert cost is not None
    assert stock_decimal is not None

    job.purchase_cost_cny = cost
    job.stock = int(stock_decimal)
    job.color = color_value
    job.tn_ved = tn_ved_value
    job.compliance_documents_json = documents
    job.cost_confirmed = True
    job.dimensions_confirmed = True
    job.stock_confirmed = True
    job.category_confirmed = True
    job.compliance_confirmed = True

    converted = convert_package_dimensions(
        length_mm=job.length_mm,
        width_mm=job.width_mm,
        height_mm=job.height_mm,
        weight_g=job.weight_g,
    )
    job.brand = NO_BRAND
    job.sanitized_title = sanitization.title
    job.sanitized_description = sanitization.description
    job.quality_report_json = sanitization.as_dict()
    job.template_file_path = rule.template_file_path
    job.image_status = "source_ready"
    job.status = "priced"

    margin_value = target_net_margin if target_net_margin is not None else job.target_net_margin
    channel_value = shipping_channel if shipping_channel is not None else job.shipping_channel
    try:
        pricing = calculate_price(
            purchase_cost_cny=cost,
            domestic_cost_cny=job.domestic_cost_cny,
            weight_g=job.weight_g,
            length_cm=converted["length_cm"],
            width_cm=converted["width_cm"],
            height_cm=converted["height_cm"],
            commission_category=rule.commission_category,
            target_net_margin=margin_value,
            shipping_channel=channel_value,
        )
    except PricingError as exc:
        _block_job(session, job, [str(exc)])

    job.price_rub = pricing["price_rub"]
    job.pricing_result_json = pricing_json_safe(pricing)
    job.price_at = datetime.utcnow()
    session.flush()

    excel_row = build_dispenser_excel_row(job, converted, images, material)
    try:
        dry_run_result = render_dry_run_workbook(rule, job, excel_row)
    except SingleSkuWorkflowError as exc:
        _block_job(session, job, exc.blockers)

    job.dry_run_result_json = dry_run_result
    job.dry_run_at = datetime.utcnow()
    job.status = "dry_run_ready"
    session.commit()
    session.refresh(job)
    return job
